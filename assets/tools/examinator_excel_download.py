import pandas as pd
import io
from js import document, Blob, URL, FileReader, Uint8Array, alert
from pyodide.ffi import create_proxy
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Constants
DOWNLOAD_BUTTON_ID = 'download-examinator-processed'
FILE_INPUT_ID = 'examinatorFile'
PROCESSING_INDICATOR_ID = 'processing-indicator-examinator'
OUTPUT_FILENAME = 'examinatoren_osiris.xlsx'
TABLE_NAME = 'ExaminatorTable'
TABLE_STYLE = 'TableStyleMedium2'

download_url = None  # Store the blob URL globally


def _reset_download_button():
    """Reset the download button to disabled state."""
    global download_url
    link = document.getElementById(DOWNLOAD_BUTTON_ID)
    link.classList.add('disabled')
    link.setAttribute('aria-disabled', 'true')
    link.style.pointerEvents = 'none'
    link.style.opacity = '0.6'
    download_url = None
    
    indicator = document.getElementById(PROCESSING_INDICATOR_ID)
    indicator.style.display = 'none'


def _enable_download_button():
    """Enable the download button."""
    link = document.getElementById(DOWNLOAD_BUTTON_ID)
    link.classList.remove('disabled')
    link.setAttribute('aria-disabled', 'false')
    link.style.pointerEvents = 'auto'
    link.style.opacity = '1'
    
    indicator = document.getElementById(PROCESSING_INDICATOR_ID)
    indicator.style.display = 'none'


def _extract_examiner_data(html_content):
    """Extract examiner data from HTML content."""
    osiris = BeautifulSoup(html_content, 'html.parser')
    examiners_rows = []

    print(f'[examinator_excel_download] Extracting examiner data...')

    for p in osiris.select('p.c51'):
        try:
            medewerker = p.text.split()[-1][1:-1]
            table = p.find_next('table')
            if not table:
                continue
                
            for course in table.select('td.c62 span.c64'):
                examiners_rows.append({
                    'Medewerker': medewerker,
                    'Cursus': course.text
                })
        except (IndexError, AttributeError) as e:
            print(f'[examinator] Error parsing examiner data: {e}')
            continue

    return examiners_rows


def _create_excel_with_table(dataframe):
    """Create an Excel file with formatted table."""
    excel_buffer = io.BytesIO()
    dataframe.to_excel(excel_buffer, index=False, sheet_name='Sheet1')
    excel_buffer.seek(0)

    print(f'[examinator_excel_download] Creating Excel file with table formatting...')
    
    # Load workbook and add table formatting
    wb = load_workbook(excel_buffer)
    ws = wb['Sheet1']
    
    # Create table
    tab = Table(displayName=TABLE_NAME, ref=ws.dimensions)
    style = TableStyleInfo(
        name=TABLE_STYLE,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    tab.tableStyleInfo = style
    ws.add_table(tab)

    # Auto-adjust column widths
    for column_cells in ws.columns:
        length = max(
            len(str(cell.value)) if cell.value is not None else 0 
            for cell in column_cells
        )
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = length + 2

    # Save to new buffer
    excel_buffer2 = io.BytesIO()
    wb.save(excel_buffer2)
    excel_buffer2.seek(0)
    
    return excel_buffer2.getvalue()


def reset_download_button(event=None):
    """Public interface for resetting download button."""
    _reset_download_button()


def trigger_download(event):
    """Handle download button click."""
    global download_url
    event.preventDefault()
    if download_url:
        a = document.createElement('a')
        a.href = download_url
        a.download = OUTPUT_FILENAME
        document.body.appendChild(a)
        a.click()
        document.body.removeChild(a)


def process_file(event):
    """Process the uploaded HTML file and create Excel output."""
    global download_url

    print(f'[examinator_excel_download] processing file...')
    
    file_input = document.getElementById(FILE_INPUT_ID)
    if not file_input.files.length:
        _reset_download_button()
        return
    
    _reset_download_button()
    file = file_input.files.item(0)
    reader = FileReader.new()

    def onload(e):
        global download_url
        try:
            html_content = e.target.result
            
            # Extract examiner data
            examiners_rows = _extract_examiner_data(html_content)
            
            if not examiners_rows:
                alert('Geen examinator data gevonden in het bestand.')
                _reset_download_button()
                return
            
            # Create DataFrame and Excel file
            df = pd.DataFrame(examiners_rows)
            excel_data = _create_excel_with_table(df)
            
            # Create downloadable blob
            uint8 = Uint8Array.new(list(excel_data))
            blob = Blob.new([uint8], {
                'type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            })
            download_url = URL.createObjectURL(blob)
            
            _enable_download_button()

            # Dispatch event to notify that processing is complete
            event = document.createEvent('Event')
            event.initEvent('excelProcessingComplete', True, True)
            document.dispatchEvent(event)

        except Exception as err:
            error_msg = f'Fout bij verwerken bestand: {err}'
            print(f'[examinator_excel_download] {error_msg}')
            alert(error_msg)
            _reset_download_button()

    reader.onload = create_proxy(onload)
    reader.readAsText(file)

# Event listeners
document.getElementById(FILE_INPUT_ID).addEventListener(
    'process-excel', create_proxy(process_file)
)
document.getElementById(DOWNLOAD_BUTTON_ID).addEventListener(
    'click', create_proxy(trigger_download)
)
document.addEventListener(
    'requestAutoSave', create_proxy(trigger_download)
)